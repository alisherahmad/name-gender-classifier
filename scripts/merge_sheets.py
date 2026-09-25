"""
Merge the raw multi-sheet FBR Active Taxpayer List (ATL) Excel export into
one clean CSV: dedupes by NTN, drops blank names, and drops rows that look
like businesses (M/S, LIMITED, INDUSTRIES, ...) or have obviously misaligned
data (a long digit string stuck inside the NAME field).

Usage:
    python scripts/merge_sheets.py --input ATL_IT.xlsx --output data/ATL_IT_merged_clean.csv
"""

import argparse
import csv
import openpyxl

BUSINESS_KEYWORDS = (
    r'M/S|PVT|LIMITED|\bLTD\b|COMPANY|ENTERPRISE|INDUSTR|ASSOCIATES|TRADERS|STORE|'
    r'CNG|TEXTILE|CONSTRUCTION|BUILDERS|MOTORS|PETROLEUM|ENGINEERING|TECHNOLOGIES|'
    r'SOLUTIONS|SERVICES|GROUP|CORPORATION|BROTHERS|SONS\b|& CO'
)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--input', required=True, help='Raw ATL .xlsx with one sheet per part')
    ap.add_argument('--output', default='data/ATL_IT_merged_clean.csv')
    args = ap.parse_args()

    import re
    biz_re = re.compile(BUSINESS_KEYWORDS, re.IGNORECASE)
    digit_re = re.compile(r'\d{5,}')

    wb = openpyxl.load_workbook(args.input, read_only=True)
    seen_ntn = set()
    written = 0

    with open(args.output, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(['SR_NO', 'NTN', 'NAME'])

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            first = True
            for row in ws.iter_rows(values_only=True):
                if first:
                    first = False
                    continue
                sr_no, ntn, name = row[0], row[1], row[2]
                if name is None or str(name).strip() == '':
                    continue
                name_clean = ' '.join(str(name).strip().split())
                ntn_clean = str(ntn).strip() if ntn is not None else ''

                if ntn_clean:
                    if ntn_clean in seen_ntn:
                        continue
                    seen_ntn.add(ntn_clean)

                if biz_re.search(name_clean) or digit_re.search(name_clean):
                    continue

                writer.writerow([sr_no, ntn_clean, name_clean])
                written += 1
            print(f'  processed {sheet_name}, running total: {written:,}')

    print(f'Done. Wrote {written:,} rows to {args.output}')


if __name__ == '__main__':
    main()
